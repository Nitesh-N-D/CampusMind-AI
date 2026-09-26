"""Multi-format ingestion: every supported format reaches the same
chunk/embed/trust pipeline, and bad files fail with a clear message instead
of a 500."""
import io
import os

import pytest
from PIL import Image, ImageDraw, ImageFont

from app.core.config import settings
from app.db import models
from app.db.database import get_db
from app.main import app

FORMAT_SAMPLES = [
    ("hostel_rules.docx", "word", "9:30 PM"),
    ("transport_schedule.xlsx", "excel", "Starting point: Tambaram; Departure time: 7:10 AM"),
    ("placement_orientation.pptx", "presentation", "minimum CGPA of 7.0"),
    ("exam_fee_schedule.csv", "csv", "Semester exam fee: Rs. 1500"),
    ("scholarship_faq.txt", "text", "15 October 2025"),
    ("library_notice.png", "image", "Library opens at 8:00 AM on weekdays"),
    ("attendance_reg_2025.pdf", "pdf", "75%"),
]


def _post(client, token, filename, content, title="Test document"):
    return client.post(
        "/api/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": (filename, content, "application/octet-stream")},
        data={"title": title, "document_type": "general_notice", "is_official": "true"},
    )


def _chunks(document_id):
    # Same test database session the API uses (conftest's override).
    session_gen = app.dependency_overrides[get_db]()
    db = next(session_gen)
    try:
        return (
            db.query(models.DocumentChunk)
            .filter(models.DocumentChunk.document_id == document_id)
            .order_by(models.DocumentChunk.chunk_index)
            .all()
        )
    finally:
        session_gen.close()


def _text_image(lines, fmt="PNG"):
    img = Image.new("RGB", (1200, 90 * len(lines) + 80), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default(size=44)
    for i, line in enumerate(lines):
        draw.text((50, 40 + i * 90), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, fmt)
    return buf.getvalue()


@pytest.mark.parametrize("filename,file_type,expected_text", FORMAT_SAMPLES)
def test_each_supported_format_is_ingested(client, college_and_admin, seed_pdf_path, filename, file_type, expected_text):
    admin_token, _ = college_and_admin
    with open(os.path.join(seed_pdf_path, filename), "rb") as f:
        resp = _post(client, admin_token, filename, f.read())
    assert resp.status_code == 200, resp.text
    doc = resp.json()
    assert doc["status"] == "ready", doc["processing_error"]
    assert doc["file_type"] == file_type
    assert doc["trust_score"] > 0

    chunks = _chunks(doc["id"])
    assert chunks, "no chunks were created"
    assert all(c.embedding for c in chunks)
    assert expected_text in " ".join(c.content for c in chunks)


def test_non_paginated_formats_cite_sections_not_pages(client, college_and_admin, seed_pdf_path):
    admin_token, _ = college_and_admin
    with open(os.path.join(seed_pdf_path, "transport_schedule.xlsx"), "rb") as f:
        doc = _post(client, admin_token, "transport_schedule.xlsx", f.read()).json()
    chunks = _chunks(doc["id"])
    assert {c.page_number for c in chunks} == {None}
    assert {c.section for c in chunks} == {"Bus routes", "Notes"}


def test_scanned_pdf_falls_back_to_ocr(client, college_and_admin):
    admin_token, _ = college_and_admin
    # An image-only PDF, like a scanned circular, has no text layer at all.
    scanned = _text_image(["Hall tickets are issued on 3 November."], fmt="PDF")
    resp = _post(client, admin_token, "scanned_circular.pdf", scanned)
    assert resp.status_code == 200
    doc = resp.json()
    assert doc["status"] == "ready", doc["processing_error"]
    assert "Hall tickets are issued" in " ".join(c.content for c in _chunks(doc["id"]))


@pytest.mark.parametrize("filename", ["notice.jpeg", "NOTICE.JPEG", "notice.jpg"])
def test_jpeg_photos_are_ingested(client, college_and_admin, filename):
    admin_token, _ = college_and_admin
    notice = _text_image(["Canteen closes at 9:00 PM on Sundays."], fmt="JPEG")
    doc = _post(client, admin_token, filename, notice).json()
    assert doc["status"] == "ready", doc["processing_error"]
    assert doc["file_type"] == "image"
    assert "Canteen closes at 9:00 PM" in " ".join(c.content for c in _chunks(doc["id"]))


def test_sideways_phone_photo_is_read(client, college_and_admin):
    admin_token, _ = college_and_admin
    upright = Image.open(io.BytesIO(_text_image(["Canteen closes at 9:00 PM on Sundays."])))
    # Phone cameras save the pixels sideways and set EXIF Orientation=6
    # ("rotate 90 degrees clockwise to display").
    sideways = upright.convert("RGB").rotate(90, expand=True)
    exif = Image.Exif()
    exif[0x0112] = 6
    buf = io.BytesIO()
    sideways.save(buf, "JPEG", exif=exif)
    doc = _post(client, admin_token, "phone_photo.jpeg", buf.getvalue()).json()
    assert doc["status"] == "ready", doc["processing_error"]
    assert "Canteen closes at 9:00 PM" in " ".join(c.content for c in _chunks(doc["id"]))


def test_pptx_cites_slides_and_reads_tables_and_notes(client, college_and_admin, seed_pdf_path):
    admin_token, _ = college_and_admin
    with open(os.path.join(seed_pdf_path, "placement_orientation.pptx"), "rb") as f:
        doc = _post(client, admin_token, "placement_orientation.pptx", f.read()).json()
    assert doc["status"] == "ready", doc["processing_error"]
    chunks = _chunks(doc["id"])
    assert {c.page_number for c in chunks} == {None}
    assert "Slide 2: Eligibility" in {c.section for c in chunks}
    text = " ".join(c.content for c in chunks)
    assert "Phase: Mock interviews; Dates: 4 August to 8 August 2025." in text  # table
    assert "Registration closes on 30 June 2025." in text  # speaker notes


def test_picture_only_slide_falls_back_to_ocr(client, college_and_admin):
    from pptx import Presentation
    from pptx.util import Inches

    admin_token, _ = college_and_admin
    # A slide that's just a pasted screenshot of a notice, no text boxes.
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    picture = io.BytesIO(_text_image(["Fee payment closes on 12 December."]))
    slide.shapes.add_picture(picture, Inches(0.5), Inches(1), width=Inches(9))
    buf = io.BytesIO()
    deck.save(buf)

    doc = _post(client, admin_token, "fee_notice.pptx", buf.getvalue()).json()
    assert doc["status"] == "ready", doc["processing_error"]
    assert "Fee payment closes" in " ".join(c.content for c in _chunks(doc["id"]))


def test_long_heading_fits_section_column(client, college_and_admin):
    from pptx import Presentation

    admin_token, _ = college_and_admin
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Revised guidelines " * 12  # ~230 characters
    slide.placeholders[1].text_frame.text = "Internal assessment carries 40 marks."
    buf = io.BytesIO()
    deck.save(buf)

    doc = _post(client, admin_token, "guidelines.pptx", buf.getvalue()).json()
    assert doc["status"] == "ready", doc["processing_error"]
    chunks = _chunks(doc["id"])
    # SQLite doesn't enforce String(120); Postgres would reject the insert.
    assert all(len(c.section) <= 120 for c in chunks)
    assert all(len(c.heading) > 120 for c in chunks)  # full heading kept for citations


@pytest.mark.parametrize(
    "filename,expected",
    [
        ("malware.exe", "Unsupported file type (.exe)"),
        ("README", "Unsupported file type"),
        ("old_rules.doc", "Save it as .docx"),
        ("old_sheet.xls", "Save it as .xlsx"),
        ("old_slides.ppt", "Save it as .pptx"),
    ],
)
def test_unsupported_types_are_rejected_with_clear_message(client, college_and_admin, filename, expected):
    admin_token, _ = college_and_admin
    resp = _post(client, admin_token, filename, b"some bytes")
    assert resp.status_code == 400
    assert expected in resp.json()["detail"]


@pytest.mark.parametrize(
    "filename,content",
    [
        ("notes.pdf", b"just some plain text, not a pdf"),
        ("rules.docx", b"plain text renamed to docx"),
        ("sheet.xlsx", b"%PDF-1.4 a pdf renamed to xlsx"),
        ("slides.pptx", b"%PDF-1.4 a pdf renamed to pptx"),
        ("photo.png", b"GIF89a not a png"),
        ("notes.txt", b"MZ\x90\x00\x03\x00\x00\x00binary"),
    ],
)
def test_renamed_files_are_rejected_before_processing(client, college_and_admin, filename, content):
    admin_token, _ = college_and_admin
    resp = _post(client, admin_token, filename, content)
    assert resp.status_code == 400
    assert "don't look like" in resp.json()["detail"]


@pytest.mark.parametrize(
    "filename,content,expected",
    [
        ("broken.pdf", b"%PDF-1.4\n garbage that is not a real pdf body", "PDF appears to be damaged"),
        ("broken.docx", b"PK\x03\x04 truncated zip", "Word file appears to be damaged"),
        ("broken.xlsx", b"PK\x03\x04 truncated zip", "Excel file appears to be damaged"),
        ("broken.pptx", b"PK\x03\x04 truncated zip", "PowerPoint file appears to be damaged"),
        ("broken.png", b"\x89PNG\r\n\x1a\n truncated", "image appears to be damaged"),
        ("blank.txt", b"   \n\n  ", "doesn't contain any readable text"),
    ],
)
def test_corrupt_files_fail_gracefully_with_clear_message(client, college_and_admin, filename, content, expected):
    admin_token, _ = college_and_admin
    resp = _post(client, admin_token, filename, content)
    assert resp.status_code == 200, resp.text
    doc = resp.json()
    assert doc["status"] == "failed"
    assert expected in doc["processing_error"]


def test_image_without_text_fails_with_clear_message(client, college_and_admin):
    admin_token, _ = college_and_admin
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, "PNG")
    doc = _post(client, admin_token, "blank.png", buf.getvalue()).json()
    assert doc["status"] == "failed"
    assert doc["processing_error"] == "No readable text was found in this image."


def test_empty_and_oversized_files_are_rejected(client, college_and_admin, monkeypatch):
    admin_token, _ = college_and_admin
    resp = _post(client, admin_token, "empty.pdf", b"")
    assert resp.status_code == 400
    assert resp.json()["detail"] == "This file is empty."

    monkeypatch.setattr(settings, "max_upload_mb", 1)
    resp = _post(client, admin_token, "big.txt", b"a" * (1024 * 1024 + 1))
    assert resp.status_code == 400
    assert "exceeds 1MB" in resp.json()["detail"]


def test_csv_with_semicolons_and_latin1_text(client, college_and_admin):
    admin_token, _ = college_and_admin
    content = "Club;Coordinator\nRobotics;Dr. José Kumar\n".encode("cp1252")
    doc = _post(client, admin_token, "clubs.csv", content).json()
    assert doc["status"] == "ready", doc["processing_error"]
    assert "Club: Robotics; Coordinator: Dr. José Kumar." in " ".join(c.content for c in _chunks(doc["id"]))


def test_students_cannot_upload_any_format(client, student_token, seed_pdf_path):
    with open(os.path.join(seed_pdf_path, "scholarship_faq.txt"), "rb") as f:
        resp = _post(client, student_token, "scholarship_faq.txt", f.read())
    assert resp.status_code == 403
