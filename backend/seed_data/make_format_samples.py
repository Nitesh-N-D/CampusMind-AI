"""Generates the non-PDF demo documents in this folder, one per supported
upload format. All content is fictional demo data for CampusMind AI and is
labelled as such inside every file.

Run from backend/:  python seed_data/make_format_samples.py
"""
import csv
import os

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
DEMO = "DEMO DATA - fictional sample document for CampusMind AI."


def make_docx():
    doc = Document()
    doc.add_heading("Hostel Rules and Regulations 2025-26", level=0)
    doc.add_paragraph(DEMO)
    doc.add_heading("Curfew", level=1)
    doc.add_paragraph(
        "First-year residents must be inside the hostel by 9:30 PM on all days. "
        "Senior residents must return by 10:30 PM. Late entry requires written permission from the warden."
    )
    doc.add_heading("Guests", level=1)
    doc.add_paragraph(
        "Visitors are allowed in the common lounge between 4:00 PM and 7:00 PM only. "
        "Visitors may not enter resident rooms."
    )
    doc.add_heading("Mess timings", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Meal"
    table.rows[0].cells[1].text = "Timing"
    for meal, timing in [
        ("Breakfast", "7:00 AM to 9:00 AM"),
        ("Lunch", "12:30 PM to 2:00 PM"),
        ("Dinner", "7:30 PM to 9:00 PM"),
    ]:
        row = table.add_row()
        row.cells[0].text = meal
        row.cells[1].text = timing
    doc.save(os.path.join(HERE, "hostel_rules.docx"))


def make_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Bus routes"
    ws.append(["Route", "Starting point", "Departure time", "Arrives at campus"])
    for row in [
        ["Route 5", "Chromepet", "7:00 AM", "8:05 AM"],
        ["Route 12", "Tambaram", "7:10 AM", "8:15 AM"],
        ["Route 18", "Velachery", "6:55 AM", "8:10 AM"],
    ]:
        ws.append(row)
    notes = wb.create_sheet("Notes")
    notes.append(["Note"])
    notes.append([DEMO])
    notes.append(["The annual transport fee is Rs. 18,000 and is payable before the start of the odd semester."])
    wb.save(os.path.join(HERE, "transport_schedule.xlsx"))


def make_pptx():
    deck = Presentation()
    title_slide = deck.slides.add_slide(deck.slide_layouts[0])
    title_slide.shapes.title.text = "Placement Orientation 2025-26"
    title_slide.placeholders[1].text = DEMO

    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Eligibility"
    body = slide.placeholders[1].text_frame
    body.text = "Students need a minimum CGPA of 7.0 to register for campus placements."
    body.add_paragraph().text = "Students with more than one standing arrear cannot register."
    slide.notes_slide.notes_text_frame.text = "Registration closes on 30 June 2025."

    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Drive schedule"
    rows = [
        ("Phase", "Dates"),
        ("Aptitude training", "1 July to 31 July 2025"),
        ("Mock interviews", "4 August to 8 August 2025"),
        ("Placement drives", "From 1 September 2025"),
    ]
    table = slide.shapes.add_table(len(rows), 2, Inches(0.8), Inches(1.8), Inches(8), Inches(2.5)).table
    for r, (phase, dates) in enumerate(rows):
        table.cell(r, 0).text = phase
        table.cell(r, 1).text = dates
    deck.save(os.path.join(HERE, "placement_orientation.pptx"))


def make_csv():
    with open(os.path.join(HERE, "exam_fee_schedule.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Programme", "Semester exam fee", "Revaluation fee per paper", "Note"])
        w.writerow(["B.E.", "Rs. 1500", "Rs. 400", DEMO])
        w.writerow(["M.E.", "Rs. 2000", "Rs. 500", ""])
        w.writerow(["MCA", "Rs. 1800", "Rs. 450", ""])


def make_txt():
    text = f"""# Scholarship FAQ 2025-26

{DEMO}

# Merit scholarship

Students with a CGPA of 8.5 or above are eligible for the merit scholarship.
The merit scholarship application deadline is 15 October 2025.

# First-generation graduate scholarship

Students who are the first in their family to attend college receive a full
tuition fee waiver. Submit the first-graduate certificate to the scholarship
cell in the admin block.
"""
    with open(os.path.join(HERE, "scholarship_faq.txt"), "w", encoding="utf-8") as f:
        f.write(text)


def make_png():
    lines = [
        "CENTRAL LIBRARY NOTICE",
        "DEMO DATA - fictional sample for CampusMind AI",
        "Library opens at 8:00 AM on weekdays.",
        "Library closes at 8:00 PM on weekdays.",
        "Late fee is Rs. 5 per day per book.",
    ]
    img = Image.new("RGB", (1400, 90 * len(lines) + 80), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default(size=44)
    for i, line in enumerate(lines):
        draw.text((50, 40 + i * 90), line, fill="black", font=font)
    img.save(os.path.join(HERE, "library_notice.png"))


if __name__ == "__main__":
    make_docx()
    make_xlsx()
    make_pptx()
    make_csv()
    make_txt()
    make_png()
    print("Wrote demo samples to", HERE)
