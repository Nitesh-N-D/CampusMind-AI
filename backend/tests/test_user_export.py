"""Admin account export: students and faculty of the admin's own college
only, as CSV or Excel, with signup date, last login, and account status -
and never a password hash."""
import csv
import io
from datetime import datetime

from openpyxl import load_workbook

from app.db import models
from app.db.database import get_db
from app.main import app

HEADERS = ["Name", "Email", "Role", "Account status", "Signed up (UTC+00:00)", "Last login (UTC+00:00)"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


def _register(client, email, name, role="student"):
    resp = client.post(
        "/api/auth/register-student",
        json={"email": email, "full_name": name, "password": "pass1234", "role": role},
    )
    assert resp.status_code == 200, resp.text


def _login(client, email):
    assert client.post("/api/auth/login", json={"email": email, "password": "pass1234"}).status_code == 200


def _export(client, token, fmt="csv", tz_offset=None):
    params = {"format": fmt}
    if tz_offset is not None:
        params["tz_offset"] = tz_offset
    return client.get("/api/admin/users/export", headers=_auth(token), params=params)


def _csv_rows(resp):
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))


def _set_user(email, **fields):
    session_gen = app.dependency_overrides[get_db]()
    db = next(session_gen)
    try:
        user = db.query(models.User).filter(models.User.email == email).one()
        for key, value in fields.items():
            setattr(user, key, value)
        db.commit()
    finally:
        session_gen.close()


def _second_college(client):
    resp = client.post(
        "/api/auth/register-college",
        json={
            "college_name": "Other Institute",
            "official_domain": "other.edu",
            "admin_email": "admin@other.edu",
            "admin_full_name": "Other Admin",
            "admin_password": "adminpass123",
        },
    )
    assert resp.status_code == 200, resp.text
    _register(client, "outsider@other.edu", "Outsider Student")
    return resp.json()["access_token"]


def _faculty(client, admin_token, domain):
    client.put("/api/admin/settings/faculty-domain", headers=_auth(admin_token),
               json={"faculty_domain": f"faculty.{domain}"})
    _register(client, f"rao@faculty.{domain}", "Prof. Rao", role="faculty")


def test_csv_lists_students_and_faculty_with_status_and_last_login(client, college_and_admin):
    admin_token, domain = college_and_admin
    _register(client, f"asha@{domain}", "Asha")
    _register(client, f"bala@{domain}", "Bala")
    _register(client, f"chitra@{domain}", "Chitra")
    _faculty(client, admin_token, domain)
    _login(client, f"asha@{domain}")
    _set_user(f"asha@{domain}", created_at=datetime(2026, 1, 5, 4, 30))
    _set_user(f"bala@{domain}", status=models.UserStatus.PENDING)
    _set_user(f"chitra@{domain}", is_active=False)

    resp = _export(client, admin_token)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert 'filename="campusmind-accounts-' in resp.headers["content-disposition"]
    rows = _csv_rows(resp)
    assert rows[0] == HEADERS
    by_email = {r[1]: r for r in rows[1:]}
    # The admin's own account isn't part of the export.
    assert set(by_email) == {f"asha@{domain}", f"bala@{domain}", f"chitra@{domain}", f"rao@faculty.{domain}"}

    asha = by_email[f"asha@{domain}"]
    assert asha[:4] == ["Asha", f"asha@{domain}", "student", "active"]
    assert asha[4] == "2026-01-05 04:30"
    assert asha[5] != "Never"
    assert by_email[f"bala@{domain}"][3] == "pending"
    assert by_email[f"bala@{domain}"][5] == "Never"  # registered, never signed in again
    assert by_email[f"chitra@{domain}"][3] == "disabled"
    assert by_email[f"rao@faculty.{domain}"][2] == "faculty"


def test_times_follow_the_admins_timezone(client, college_and_admin):
    admin_token, domain = college_and_admin
    _register(client, f"asha@{domain}", "Asha")
    _set_user(f"asha@{domain}", created_at=datetime(2026, 1, 5, 20, 0))
    rows = _csv_rows(_export(client, admin_token, tz_offset=330))
    assert rows[0][4] == "Signed up (UTC+05:30)"
    assert rows[1][4] == "2026-01-06 01:30"


def test_xlsx_export_opens_with_the_same_columns(client, college_and_admin):
    admin_token, domain = college_and_admin
    _register(client, f"asha@{domain}", "Asha")
    _login(client, f"asha@{domain}")

    resp = _export(client, admin_token, "xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sheet = load_workbook(io.BytesIO(resp.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert rows[1][:4] == ("Asha", f"asha@{domain}", "student", "active")
    assert isinstance(rows[1][4], datetime) and isinstance(rows[1][5], datetime)


def test_export_never_contains_password_hashes(client, college_and_admin):
    admin_token, domain = college_and_admin
    _register(client, f"asha@{domain}", "Asha")
    for fmt in ("csv", "xlsx"):
        body = _export(client, admin_token, fmt).content
        assert b"$2b$" not in body and b"pass1234" not in body
    header = _csv_rows(_export(client, admin_token))[0]
    assert not any("password" in h.lower() or "hash" in h.lower() for h in header)


def test_export_is_scoped_to_the_admins_own_college(client, college_and_admin):
    admin_token, domain = college_and_admin
    _register(client, f"asha@{domain}", "Asha")
    other_admin = _second_college(client)

    for token, mine, theirs in [(admin_token, f"asha@{domain}", "outsider@other.edu"),
                                (other_admin, "outsider@other.edu", f"asha@{domain}")]:
        csv_body = _export(client, token).content.decode("utf-8-sig")
        assert mine in csv_body and theirs not in csv_body
        xlsx_values = [c for row in load_workbook(io.BytesIO(_export(client, token, "xlsx").content)).active
                       .iter_rows(values_only=True) for c in row]
        assert mine in xlsx_values and theirs not in xlsx_values


def test_formula_like_names_are_neutralised(client, college_and_admin):
    admin_token, domain = college_and_admin
    _register(client, f"asha@{domain}", "=HYPERLINK(\"http://evil\")")
    rows = _csv_rows(_export(client, admin_token))
    assert rows[1][0].startswith("'=")


def test_only_admins_can_export(client, college_and_admin, student_token):
    admin_token, domain = college_and_admin
    _faculty(client, admin_token, domain)
    faculty_token = client.post(
        "/api/auth/login", json={"email": f"rao@faculty.{domain}", "password": "pass1234"}
    ).json()["access_token"]
    assert client.get("/api/admin/users/export").status_code == 401
    assert _export(client, student_token).status_code == 403
    assert _export(client, faculty_token).status_code == 403
    assert _export(client, admin_token, "pdf").status_code == 422
